"""Формирование доверенности по образцу doverennost_primer."""
from __future__ import annotations

from datetime import date, timedelta

from django.utils import timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from apps.users.models import UserProfile

MONTHS_GENITIVE = (
    "",
    "января",
    "февраля",
    "марта",
    "апреля",
    "мая",
    "июня",
    "июля",
    "августа",
    "сентября",
    "октября",
    "ноября",
    "декабря",
)

# (текст, жирный)
Segment = tuple[str, bool]


def format_date_ru(d: date) -> str:
    return f"{d.day} {MONTHS_GENITIVE[d.month]} {d.year} года"


def _add_rich_paragraph(
    doc: Document,
    segments: list[Segment],
    *,
    align=None,
    size: int = 12,
):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    for text, bold in segments:
        if not text:
            continue
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
    return p


def _burial_registration_label(burial) -> str:
    if burial.global_id:
        return str(burial.global_id)
    if burial.grave_number:
        return burial.grave_number
    return str(burial.pk)


def _burial_certificate_hint(burial) -> str:
    if burial.grave_number:
        return f"удостоверение о захоронении, участок {burial.grave_number}"
    return "удостоверение о захоронении (номер уточняется в администрации кладбища)"


def build_power_of_attorney_docx(
    *,
    trustor: UserProfile,
    attorney: UserProfile,
    burial,
    issue_date: date | None = None,
    valid_until: date | None = None,
) -> Document:
    """DOCX по шаблону: доверитель (заказчик) уполномочивает поверенного (исполнителя)."""
    issue_date = issue_date or timezone.localdate()
    valid_until = valid_until or (issue_date + timedelta(days=365))
    city = (trustor.poa_city or "г. Москва").strip()
    cemetery = burial.cemetery
    cemetery_name = cemetery.name if cemetery else "—"
    reg_no = _burial_registration_label(burial)
    cert_hint = _burial_certificate_hint(burial)

    trustor_name = trustor.full_name
    attorney_name = attorney.full_name
    trustor_passport = trustor.passport_line
    attorney_passport = attorney.passport_line
    issue_date_s = format_date_ru(issue_date)
    valid_until_s = format_date_ru(valid_until)

    doc = Document()

    _add_rich_paragraph(
        doc,
        [("ДОВЕРЕННОСТЬ", True)],
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=14,
    )
    _add_rich_paragraph(
        doc,
        [(f"{city} ", False), (issue_date_s, True)],
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    _add_rich_paragraph(
        doc,
        [
            ("Я, гражданин Российской Федерации, ", False),
            (trustor_name, True),
            (" ", False),
            (trustor_passport, True),
            (", являясь ответственным лицом за захоронение рег. № ", False),
            (reg_no, True),
            (", расположенного на кладбище «", False),
            (cemetery_name, True),
            ("», что подтверждается ", False),
            (cert_hint, True),
            (", далее – «", False),
            ("Захоронение", True),
            ("», уполномочиваю гражданина ", False),
            (attorney_name, True),
            (", ", False),
            (attorney_passport, True),
            (", далее ", False),
            ("Поверенный", True),
            (", ", False),
        ],
    )

    _add_rich_paragraph(
        doc,
        [
            ("представлять мои интересы во всех компетентных учреждениях и организациях (в том ", False),
            ("числе, но не ограничиваясь, в Администрации кладбища, АО «Честный Агент»), ", False),
            ("связанные с обустройством ", False),
            ("Захоронения", True),
            (", для чего предоставляю ", False),
            ("Поверенному", True),
            (" право: ", False),
        ],
    )

    _add_rich_paragraph(
        doc,
        [
            (
                "заключать договоры/дополнительные соглашения, подписывать заказы и акты приема-"
                "передачи к ранее заключенным договорам на оказание различного вида услуг по "
                "обустройству, уборке, уходу за захоронением, проведением монтажных/демонтажных "
                "работ на месте захоронения, такие как: установка/замена памятника, надгробия, "
                "ограждения; заказ граверных работ на памятнике (надпись, изображение); заказ товаров и "
                "услуг, связанных с благоустройством участка, на котором находится захоронение "
                "(приобретение и покраска ограды, установка цоколя, цветника, лавки, стола, вазы, укладка "
                "плитки, отсыпка участка песком и/или гранитным щебнем); получать все необходимые "
                "документы; ставить свою подпись на всех необходимых документах, в том числе на "
                "договоре/дополнительном соглашении и т.п.; производить все необходимые платежи и "
                "совершать все иные действия, связанные с выполнением настоящего поручения.",
                False,
            ),
        ],
    )

    _add_rich_paragraph(
        doc,
        [("Доверенность выдана сроком до ", False), (valid_until_s, True)],
    )
    _add_rich_paragraph(
        doc,
        [("Доверитель: ", True), (trustor_name, True)],
    )

    return doc
