from __future__ import annotations

from datetime import timedelta
from pathlib import Path

from django.conf import settings
from django.utils import timezone
from docx import Document

from .models import (
    CareSubscription,
    Notification,
    NotificationType,
    Order,
    OrderStatus,
    PowerOfAttorneyDocument,
    SubscriptionPeriod,
)


VALID_TRANSITIONS = {
    OrderStatus.CREATED: {OrderStatus.IN_PROGRESS, OrderStatus.CANCELLED},
    OrderStatus.IN_PROGRESS: {OrderStatus.DONE, OrderStatus.CANCELLED},
    OrderStatus.DONE: {OrderStatus.CONFIRMED},
    OrderStatus.CONFIRMED: set(),
    OrderStatus.CANCELLED: set(),
}


def transition_order(order: Order, *, target_status: str, actor_id: int | None = None) -> Order:
    allowed = VALID_TRANSITIONS.get(order.status, set())
    if target_status not in allowed:
        raise ValueError(f"Недопустимый переход статуса {order.status} -> {target_status}")
    order.status = target_status
    order.save(update_fields=["status", "updated_at"])

    _create_order_notification(order, actor_id=actor_id, status=target_status)
    return order


def generate_power_of_attorney(order: Order) -> PowerOfAttorneyDocument:
    doc = Document()
    doc.add_heading("Доверенность на уход за захоронением", level=1)
    doc.add_paragraph(f"Дата: {timezone.localdate().isoformat()}")
    doc.add_paragraph(f"Заказ № {order.id}")
    doc.add_paragraph(f"Заказчик: {order.customer.username}")
    doc.add_paragraph(f"Исполнитель: {order.executor.username if order.executor else 'не назначен'}")
    doc.add_paragraph(f"Кладбище: {order.burial.cemetery.name}")
    doc.add_paragraph(f"Участок: {order.burial.grave_number or 'не указан'}")
    doc.add_paragraph(
        "Настоящим Заказчик поручает Исполнителю выполнение работ по уходу "
        "за указанным захоронением в рамках оформленного заказа."
    )

    rel_dir = Path("documents/power_of_attorney")
    abs_dir = Path(settings.MEDIA_ROOT) / rel_dir
    abs_dir.mkdir(parents=True, exist_ok=True)
    filename = f"order_{order.id}_{timezone.now().strftime('%Y%m%d_%H%M%S')}.docx"
    abs_path = abs_dir / filename
    doc.save(str(abs_path))

    obj, _ = PowerOfAttorneyDocument.objects.update_or_create(
        order=order,
        defaults={"file": str(rel_dir / filename)},
    )
    return obj


def _create_order_notification(order: Order, *, actor_id: int | None, status: str) -> None:
    title = f"Заказ #{order.id}: статус {order.get_status_display()}"
    body = "Статус заказа изменен."
    payload = {"order_id": order.id, "status": status, "actor_id": actor_id}
    user_ids = {order.customer_id}
    if order.executor_id:
        user_ids.add(order.executor_id)
    for user_id in user_ids:
        Notification.objects.create(
            user_id=user_id,
            type=NotificationType.ORDER_EVENT,
            title=title,
            body=body,
            payload=payload,
        )


def build_next_subscription_date(period: str, base_dt=None):
    now = base_dt or timezone.now()
    if period == SubscriptionPeriod.QUARTERLY:
        return now + timedelta(days=90)
    return now + timedelta(days=30)


def create_order_from_subscription(sub: CareSubscription) -> Order:
    order = Order.objects.create(
        customer=sub.customer,
        burial=sub.burial,
        service_type=sub.service_type,
        status=OrderStatus.CREATED,
        description="Автосозданный заказ по подписке",
    )
    sub.next_run_at = build_next_subscription_date(sub.period, base_dt=sub.next_run_at)
    sub.save(update_fields=["next_run_at"])
    return order
